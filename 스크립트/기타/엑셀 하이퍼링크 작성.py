
import os
from pathlib import Path, PureWindowsPath
from tkinter import Tk
from tkinter.filedialog import askopenfilename
from openpyxl import load_workbook
from copy import copy


# ==========================================
# 1. 사용자 설정
# ==========================================

HEADER_ROW = 1
OUTPUT_NAME = "out.xls"


# ==========================================
# 2. 엑셀 파일 선택
# ==========================================

root = Tk()
root.withdraw()

INPUT_FILE = askopenfilename(
    title="도면목록 엑셀 파일 선택",
    filetypes=[("Excel 파일", "*.xlsx *.xlsm")]
)

root.destroy()

if not INPUT_FILE:
    raise SystemExit("파일 선택이 취소되었습니다.")

input_path = Path(INPUT_FILE)
input_folder = input_path.parent

# XLSM 파일은 매크로를 보존하여 저장
is_macro = input_path.suffix.lower() == ".xlsm"

output_ext = ".xlsm" if is_macro else ".xlsx"
OUTPUT_FILE = input_folder / (
    Path(OUTPUT_NAME).stem + output_ext
)


# ==========================================
# 3. DWG 파일 재귀 탐색
# ==========================================

print("=" * 55)
print("DWG 파일 검색 시작")
print(f"검색 폴더: {input_folder}")
print("=" * 55)

# 파일명 -> 절대경로
dwg_index = {}

for current_dir, dirs, files in os.walk(input_folder):

    for filename in files:

        if not filename.lower().endswith(".dwg"):
            continue

        full_path = Path(current_dir) / filename

        # 파일명은 고유하므로 딕셔너리에 바로 등록
        dwg_index[filename.casefold()] = full_path.resolve()

print(f"DWG 검색 완료: {len(dwg_index)}개")


# ==========================================
# 4. DWG 경로 검색 함수
# ==========================================

def find_dwg(old_path):
    """
    기존 Windows 경로에서 파일명만 추출한 뒤
    실제 DWG 파일의 절대경로 반환
    """

    filename = PureWindowsPath(
        str(old_path).strip().replace("/", "\\")
    ).name.casefold()

    return dwg_index.get(filename)


# ==========================================
# 5. 엑셀 불러오기
# ==========================================

wb = load_workbook(
    INPUT_FILE,
    keep_vba=is_macro
)

total = 0
linked = 0
missing = 0

print("\n엑셀 도면경로 수정 시작\n")


# ==========================================
# 6. 모든 시트 처리
# ==========================================

for ws in wb.worksheets:

    # 헤더명으로 열 자동 인식
    headers = {
        str(cell.value).strip(): cell.column
        for cell in ws[HEADER_ROW]
        if cell.value is not None
    }

    if not {"도면경로", "주제목"}.issubset(headers):
        print(f"[SKIP] {ws.title}: 필수 열 없음")
        continue

    path_col = headers["도면경로"]
    title_col = headers["주제목"]

    print(f"[시트 처리] {ws.title}")

    for row in range(HEADER_ROW + 1, ws.max_row + 1):

        path_cell = ws.cell(row, path_col)
        title_cell = ws.cell(row, title_col)

        # 경로 또는 주제목이 비어 있으면 건너뛰기
        if not path_cell.value or not title_cell.value:
            continue

        old_path = str(path_cell.value).strip()

        total += 1

        # 실제 DWG 파일 검색
        new_path = find_dwg(old_path)

        if new_path is None:
            missing += 1
            print(f"[파일 없음] {old_path}")
            continue

        # ----------------------------------
        # 도면경로 수정
        # ----------------------------------

        path_cell.value = str(new_path)

        # ----------------------------------
        # 주제목 하이퍼링크 생성
        # ----------------------------------

        title_cell.hyperlink = str(new_path)

        # 기존 글꼴 서식 유지
        font = copy(title_cell.font)
        font.color = "0563C1"
        font.underline = "single"

        title_cell.font = font

        linked += 1

        print(
            f"[성공] {new_path.name}\n"
            f"       {new_path}"
        )


# ==========================================
# 7. 결과 저장
# ==========================================

wb.save(OUTPUT_FILE)

print("\n" + "=" * 55)
print("도면목록 처리 완료")
print("=" * 55)
print(f"검색된 DWG 파일 : {len(dwg_index)}개")
print(f"전체 도면       : {total}개")
print(f"링크 생성 성공  : {linked}개")
print(f"파일 없음       : {missing}개")
print(f"저장 경로       : {OUTPUT_FILE}")
print("=" * 55)
