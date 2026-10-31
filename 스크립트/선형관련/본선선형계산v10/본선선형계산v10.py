"""
본선 선형 계산 v10 (리팩토링)
================================
- 원본: 본선선형계산v9.py
- 주요 개선:
  1. P0 버그 수정 (coord NameError, unit 전역참조, get_station_coordinates 반환 누락, pop(0) 원본변형)
  2. 전역 변수 제거 → AppContext 클래스 캡슐화
  3. CurveParams dataclass 도입 (매직 인덱스 제거)
  4. 하드코딩 경로 → pathlib + 설정 상수
  5. 예외 처리 세분화, 로깅 도입
  6. 중복 함수 제거 (calculate_distance, format_distance)
  7. 매직 넘버 상수화
"""

from __future__ import annotations

import csv
import logging
import math
import tempfile
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional

import ezdxf
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd  # noqa: F401 (향후 확장용, 미사용 시 제거)
import tkinter as tk
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from shapely.geometry import LineString, Point

from tangent_calculator import (
    TangentFields,
    create_tangent_calculator,
)
from coordinate_calculator import (
    CoordinateResult,
    CurveGeometry,
    build_geometry_from_curve,
    create_coordinate_calculator,
)
from station_builder import build_station_group
from report_writer import write_report, save_with_retry
from dxf_writer import create_dxf
from 스크립트.선형관련.본선선형계산v10.constant_config import Config, M_VALUE_TABLE
from 스크립트.선형관련.본선선형계산v10.math_util import calculate_bearing_n, calculate_distance, format_distance

# ============================================================
# 로깅 설정
# ============================================================
logging.basicConfig(
    level=logging.INFO,
    format="[%(levelname)s] %(message)s",
)
logger = logging.getLogger(__name__)


# ============================================================
# 데이터 클래스
# ============================================================
@dataclass
class CurveParams:
    """IP별 선형 계산 결과 (매직 인덱스 제거)"""

    ip_no: int
    ip_xy: tuple[float, float]
    curve_type: str                     # 'simplecurve' | 'cubic'
    direction: str                      # 'LEFT CURVE' | 'RIGHT CURVE'
    h1: float                           # 방위각 1
    h2: float                           # 방위각 2
    ia_degree: float                    # 교각 (도)
    center_xy: tuple[float, float]
    radius: float
    ip_distance1: float
    ip_distance2: float
    m: float
    z: float
    v: float
    params: list                        # simplecurve_parameter 또는 cubic_parameter
    o_pc_bearing: float
    sp_pc_bearing: float
    theta_pc: float
    ip_sta: float

    # 측점 (curve_type에 따라 다름)
    sp_sta: Optional[float] = None
    pc_sta: Optional[float] = None
    cp_sta: Optional[float] = None
    ps_sta: Optional[float] = None
    bc_sta: Optional[float] = None
    ec_sta: Optional[float] = None

    # 좌표
    sp_xy: Optional[tuple[float, float]] = None
    pc_xy: Optional[tuple[float, float]] = None
    cp_xy: Optional[tuple[float, float]] = None
    ps_xy: Optional[tuple[float, float]] = None
    bc_xy: Optional[tuple[float, float]] = None
    ec_xy: Optional[tuple[float, float]] = None

    # 편의 프로퍼티
    @property
    def tl(self) -> float:
        if self.curve_type == "simplecurve":
            return self.params[0]
        return self.params[10]

    @property
    def cl(self) -> float:
        if self.curve_type == "simplecurve":
            return self.params[1]
        return self.params[12]

    @property
    def spiral_length(self) -> float:
        """완화곡선 길이 L (단곡선이면 0)"""
        return 0.0 if self.curve_type == "simplecurve" else self.params[3]


@dataclass
class StationRow:
    """측점 1개 행 데이터"""

    label: str
    sta: float
    station_offset: float
    coord: tuple[float, float]
    shiftx: float
    shifty: float
    ta: float                # 접선방위각
    chord_length: float      # 사거리
    q: float                 # 편기각
    t: float                 # 접선각
    azimuth: float           # 방위각


@dataclass
class AppContext:
    """main_gui 전역 상태 캡슐화"""

    linestring: Optional[LineString] = None
    curves: list[CurveParams] = field(default_factory=list)
    station_groups: list[list[list[StationRow]]] = field(default_factory=list)
    polycurve: list[tuple[float, float]] = field(default_factory=list)







def calculate_inner_angle(v1: tuple, v2: tuple) -> float:
    a1 = math.atan2(v1[1], v1[0])
    a2 = math.atan2(v2[1], v2[0])
    diff = abs(a2 - a1)
    if diff > math.pi:
        diff = 2 * math.pi - diff
    return math.degrees(diff)


def calculate_angles(line_string: LineString) -> list[float]:
    angles = []
    coords = line_string.coords
    for i in range(1, len(coords) - 1):
        v1 = (coords[i][0] - coords[i - 1][0], coords[i][1] - coords[i - 1][1])
        v2 = (coords[i + 1][0] - coords[i][0], coords[i + 1][1] - coords[i][1])
        angles.append(calculate_inner_angle(v1, v2))
    return angles


def find_direction(radius: float, h1: float, h2: float) -> int:
    if radius == 0:
        return 0
    return 1 if math.sin(math.radians(h2 - h1)) >= 0 else -1


def generate_integers(start: float, end: float, inc: int) -> list[int]:
    """start~end 사이 inc 배수 정수 리스트 (음수/역방향 지원)"""
    if inc <= 0:
        raise ValueError(f"증분은 양수여야 합니다: {inc}")
    if start > end:
        start, end = end, start
    start_rounded = math.ceil(start / inc) * inc
    end_rounded = math.floor(end / inc) * inc
    return list(range(int(start_rounded), int(end_rounded) + inc, inc))


# ============================================================
# 각도/거리 포맷
# ============================================================
def degrees_to_dms(degrees: float) -> str:
    """도분초 문자열 (부호 유지)"""
    sign = "-" if degrees < 0 else ""
    d = abs(degrees)
    deg = int(d)
    minutes = int((d - deg) * 60)
    seconds = (d - deg - minutes / 60) * 3600
    return f"{sign}{deg}도  {minutes}분  {seconds:.2f}초"







def format_station(sta: float) -> str:
    km = int(sta // 1000)
    meter = sta - km * 1000
    return f"{km}+{meter:06.2f}"


def format_bearing(angle: float) -> str:
    angle = angle % 360
    total_seconds = round(angle * 3600) % (360 * 3600)
    deg, rem = divmod(total_seconds, 3600)
    minute, second = divmod(rem, 60)
    return f'N {deg:03d}° {minute:02d}\' {second:02d}"'





def read_alignment_file() -> list[list[str]]:
    """파일 선택 및 읽기 (UTF-8 → EUC-KR 폴백)"""
    from tkinter import filedialog

    file_path = filedialog.askopenfilename(
        defaultextension=".txt",
        filetypes=[("txt files", "*.txt"), ("All files", "*.*")],
    )
    if not file_path:
        return []
    logger.info("선택 파일: %s", file_path)

    for encoding in ("utf-8", "euc-kr"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return list(csv.reader(f))
        except UnicodeDecodeError:
            logger.warning("%s 인코딩 실패, 다음 시도...", encoding)
    logger.error("파일을 읽을 수 없습니다: %s", file_path)
    return []


def parse_ip_lines(lines: list[list[str]]) -> list[list[float]]:
    """IP 데이터 파싱 (예외 처리 강화)"""
    alignment = []
    for line in lines:
        try:
            ip_y = float(line[0])
            ip_x = float(line[1])
            radius = float(line[2]) if len(line) > 2 else 0.0
        except (IndexError, ValueError) as e:
            logger.warning("IP 파싱 실패 (line=%s): %s", line, e)
            continue
        alignment.append([ip_x, ip_y, radius])
    return alignment


def create_linestring(alignment: list[list[float]]) -> LineString:
    return LineString([(p[0], p[1]) for p in alignment])


# ============================================================
# 캔트/M/Z/V 계산
# ============================================================
def get_cant_limits(v: float, track_type: str) -> tuple[int, int]:
    """설계속도·도상별 (최대설정캔트, 최대부족캔트)"""
    if track_type == "자갈도상":
        if v > 400:
            raise ValueError(f"지원 범위 초과: V={v}")
        if 350 < v <= 400:
            return 180, 130
        if 250 < v <= 350:
            return 160, 80
        return 160, 100
    if track_type == "콘크리트도상":
        if v > 400:
            raise ValueError(f"지원 범위 초과: V={v}")
        return 180, 130
    raise ValueError(f"잘못된 도상 유형: {track_type}")


def calculate_v(cm: int, cd: int, radius: float, max_v: int) -> int:

    v = int(math.sqrt((cm + cd) * radius / Config.CANT_COEFFICIENT))
    return min(v, max_v)


def calculate_parameters(
    radius_list: list[float],
    max_v: int = Config.DEFAULT_MAX_V,
    track_type: str = "자갈도상",
) -> list[list[float]]:
    """곡선별 [M, Z, V] 계산"""
    parameters = []
    for radius in radius_list:
        cm, cd = get_cant_limits(max_v, track_type)
        v = round(calculate_v(cm, cd, radius, max_v) / Config.V_ROUNDING) * Config.V_ROUNDING
        m = int(M_VALUE_TABLE.get(v, Config.M_FALLBACK_COEFFICIENT * v))
        m = round(m / Config.M_ROUNDING) * Config.M_ROUNDING
        z = min(Config.CANT_COEFFICIENT * v ** 2 / radius, cm)
        z = round(z / Config.Z_ROUNDING) * Config.Z_ROUNDING
        parameters.append([m, z, v])
    return parameters


def dictionary_cubicparabola(v: int) -> float:
    """완화곡선 미삽입 최소 곡선반경"""
    not_spiral_min_radius = {250: 24000, 200: 12000, 150: 5000, 120: 2500, 70: 600}
    cdlim = {400: 20, 350: 25, 300: 27, 250: 32, 200: 40, 150: 57, 120: 69, 100: 83, 70: 100}

    if v in not_spiral_min_radius:
        return not_spiral_min_radius[v]

    # cdlim 보간
    if v in cdlim:
        cdlim_val = cdlim[v]
    else:
        keys = sorted(cdlim)
        if v < keys[0] or v > keys[-1]:
            raise ValueError(f"속도 {v} 범위 초과 ({keys[0]}~{keys[-1]})")
        cdlim_val = None
        for i in range(len(keys) - 1):
            if keys[i] <= v <= keys[i + 1]:
                x1, y1 = keys[i], cdlim[keys[i]]
                x2, y2 = keys[i + 1], cdlim[keys[i + 1]]
                cdlim_val = y1 + (v - x1) * (y2 - y1) / (x2 - x1)
                break

    return Config.CANT_COEFFICIENT * (v ** 2 / cdlim_val)


# ============================================================
# 선형 계산 (수식 원본 유지)
# ============================================================
def calculate_theta_pc(m: float, z: float, rc: float) -> float:
    return math.atan(m * (z * 0.001) / (2 * rc))


def calculate_cubic_parabola_parameter(
    m: float, z: float, direction: int, theta_pc: float, rc: float, ia_rad: float
) -> list[float]:
    x1 = m * (z * 0.001)
    x2 = x1 - rc * math.sin(theta_pc)
    w13 = rc * direction
    L = x1 * (1 + (math.tan(theta_pc) ** 2) / 10)
    Y = (x1 ** 2) / (6 * rc)
    w15 = (x1 ** 2) / (6 * w13)
    F = Y - rc * (1 - math.cos(theta_pc))
    S = F / math.cos(ia_rad / 2)
    K = F * math.tan(ia_rad / 2)
    W = (rc + F) * math.tan(ia_rad / 2)
    TL = x2 + W
    Lc = rc * (ia_rad - 2 * theta_pc)
    CL = Lc + 2 * L
    SL = rc * (1 / math.cos(ia_rad / 2) - 1) + S
    RIA = ia_rad - 2 * theta_pc
    C = math.atan(Y / x1)
    XB = math.pi / 2 - theta_pc
    B = math.pi / 2 - C - XB
    return [x1, x2, w13, L, Y, w15, F, S, K, W, TL, Lc, CL, SL, RIA, C, B]


def calculate_simplecurve_parameter(direction: int, rc: float, ia_rad: float) -> list[float]:
    TL = rc * math.tan(ia_rad / 2)
    CL = rc * ia_rad
    SL = rc * (1 / math.cos(ia_rad / 2) - 1)
    S = rc * (1 - math.cos(ia_rad / 2))
    W13 = rc * direction
    return [TL, CL, SL, S, W13]


def calculate_sp_pc_bearing(w13: float, v10: float, x10: float) -> float:
    result = v10 - x10 if w13 < 0 else v10 + x10
    return result + 360 if result < 0 else (result - 360 if result > 360 else result)


def calculate_o_pc_bearing(v10: float, theta_pc: float, w13: float) -> float:
    x10 = math.degrees(theta_pc)
    if w13 > 0:
        r = v10 + x10 - 90
        return r + 360 if r < 0 else r
    r = v10 - x10 + 90
    return r - 360 if r > 360 else r


def calculate_cp_xy(k21, n21, w13, w18, p10, v11) -> tuple[float, float]:
    a1 = w13 * math.cos(math.radians(w18 + 90))
    a2 = w13 * math.cos(math.radians(w18 - 90 + (180 * p10) / (math.pi * w13)))
    b1 = w13 * math.sin(math.radians(w18 + 90))
    b2 = w13 * math.sin(math.radians(w18 - 90 + (180 * p10) / (math.pi * w13)))
    return k21 + a1 + a2, n21 + b1 + b2


def calculate_spiral_coordinates(
    bp_xy, ip_xy, ep_xy, h1, h2, tl, x1, y, sp_pc_bearing, rc, lc, w13
):
    sp_xy = (
        math.cos(math.radians(h1 + 180)) * tl + ip_xy[0],
        math.sin(math.radians(h1 + 180)) * tl + ip_xy[1],
    )
    pc_xy = (
        sp_xy[0] + x1 * math.cos(math.radians(h1)) + y * math.cos(math.radians(h1 + 90)),
        sp_xy[1] + x1 * math.sin(math.radians(h1)) + y * math.sin(math.radians(h1 + 90)),
    )
    cp_xy = calculate_cp_xy(pc_xy[0], pc_xy[1], w13, sp_pc_bearing, lc, h2)
    ps_xy = (
        math.cos(math.radians(h2)) * tl + ip_xy[0],
        math.sin(math.radians(h2)) * tl + ip_xy[1],
    )
    return sp_xy, pc_xy, cp_xy, ps_xy


def calculate_simplecurve_coordinates(v10, v11, tl, ip) -> list[tuple[float, float]]:
    bc = (
        math.cos(math.radians(v10 + 180)) * tl + ip[0],
        math.sin(math.radians(v10 + 180)) * tl + ip[1],
    )
    ec = (
        math.cos(math.radians(v11)) * tl + ip[0],
        math.sin(math.radians(v11)) * tl + ip[1],
    )
    return [bc, ec]


# ============================================================
# 측거/이정량
# ============================================================
def calculate_station_offset(
    sta: float, sp: float, pc: float, cp: float, ps: float, bp: float
) -> float:
    s, sp_i, pc_i, cp_i, ps_i, bp_i = map(int, (sta, sp, pc, cp, ps, bp))
    if s <= sp_i:
        return sta - bp
    if s < pc_i:
        return sta - sp
    if s <= cp_i:
        return sta - pc
    if s <= ps_i:
        return ps - sta
    return sta - ps


def cal_spiral_shiftx(sta, sp, pc, cp, ps, offset, w13, length) -> float:
    s, sp_i, pc_i, cp_i, ps_i = map(int, (sta, sp, pc, cp, ps))
    if sp_i >= s:
        return 0.0
    if pc_i > s or ps_i >= s:
        return offset - offset ** 5 / (40 * w13 ** 2 * length ** 2)
    return 0.0


def cal_spiral_shifty(sta, sp, pc, cp, ps, shiftx, w13, x1) -> float:
    s, sp_i, pc_i, cp_i, ps_i = map(int, (sta, sp, pc, cp, ps))
    if sp_i >= s:
        return 0.0
    if pc_i > s or ps_i >= s:
        return shiftx ** 3 / (6 * w13 * x1)
    return 0.0


def get_station_label(
    curve_type: str, station: float,
    sp: float, pc: float, cp: float, ps: float, ep: float,
    eps: float = 1e-6,
) -> str:
    checks = [
        (sp, "BC" if curve_type == "simplecurve" else "SP"),
        (pc, "BC" if curve_type == "simplecurve" else "PC"),
        (cp, "EC" if curve_type == "simplecurve" else "CP"),
        (ps, "EC" if curve_type == "simplecurve" else "PS"),
        (ep, "EP"),
    ]
    for target, label in checks:
        if abs(station - target) < eps:
            return label
    return ""


# ============================================================
# 메인 선형 계산
# ============================================================
def calculate_curve(
    linestring: LineString,
    radius_list: list[float],
    angles: list[float],
    parameters: list[list[float]],
    unit: int = Config.DEFAULT_UNIT,
    start_sta: float = Config.DEFAULT_START_STA,
) -> tuple[list[list[list[StationRow]]], list[CurveParams], list[tuple[float, float]]]:
    """모든 IP의 선형 계산"""
    station_groups: list[list[list[StationRow]]] = []
    curves: list[CurveParams] = []
    polycurve: list[tuple[float, float]] = []

    tl_list: list[float] = []
    end_sta_list: list[float] = []
    end_xy_list: list[tuple[float, float]] = []

    bp_sta = start_sta
    coords = linestring.coords

    for i in range(len(coords) - 2):
        ia_rad = math.radians(angles[i])
        ia_deg = angles[i]

        bp_xy = (coords[i][0], coords[i][1])
        ip_xy = (coords[i + 1][0], coords[i + 1][1])
        ep_xy = (coords[i + 2][0], coords[i + 2][1])

        h1 = calculate_bearing_n(*bp_xy, *ip_xy)
        h2 = calculate_bearing_n(*ip_xy, *ep_xy)
        ip_dist1 = calculate_distance(*bp_xy, *ip_xy)
        ip_dist2 = calculate_distance(*ip_xy, *ep_xy)

        rc = radius_list[i]
        m, z, v = parameters[i]
        not_spiral_min = dictionary_cubicparabola(v)

        curve_type = "simplecurve" if rc >= not_spiral_min else "cubic"
        direction = find_direction(rc, h1, h2)
        curvedir = "RIGHT CURVE" if direction == 1 else "LEFT CURVE"

        is_last_ip = (i == len(coords) - 3)

        # ---------- 곡선 파라미터 계산 ----------
        if curve_type == "simplecurve":
            theta_pc = 0.0
            sc_params = calculate_simplecurve_parameter(direction, rc, ia_rad)
            tl, cl, sl, s_val, w13 = sc_params

            o_pc_bearing = calculate_o_pc_bearing(h1, theta_pc, w13)
            sp_pc_bearing = h1
            bc_xy, ec_xy = calculate_simplecurve_coordinates(h1, h2, tl, ip_xy)
            center_xy = (
                math.cos(math.radians(o_pc_bearing + 180)) * rc + bc_xy[0],
                math.sin(math.radians(o_pc_bearing + 180)) * rc + bc_xy[1],
            )

            if i == 0:
                ip_sta = bp_sta + ip_dist1
                bc_sta = ip_sta - tl
                ec_sta = bc_sta + cl
            else:
                bp_sta = end_sta_list[i - 1]
                ip_sta = bp_sta + ip_dist1 - tl_list[i - 1]
                bc_sta = ip_sta - tl
                ec_sta = bc_sta + cl

            end_sta_list.append(ec_sta)
            end_xy_list.append(ec_xy)

            curves.append(CurveParams(
                ip_no=i + 1, ip_xy=ip_xy, curve_type=curve_type, direction=curvedir,
                h1=h1, h2=h2, ia_degree=ia_deg, center_xy=center_xy, radius=rc,
                ip_distance1=ip_dist1, ip_distance2=ip_dist2, m=m, z=z, v=v,
                params=sc_params, o_pc_bearing=o_pc_bearing,
                sp_pc_bearing=sp_pc_bearing, theta_pc=theta_pc, ip_sta=ip_sta,
                bc_sta=bc_sta, ec_sta=ec_sta, bc_xy=bc_xy, ec_xy=ec_xy,
            ))

        else:  # cubic
            theta_pc = calculate_theta_pc(m, z, rc)
            cb_params = calculate_cubic_parabola_parameter(m, z, direction, theta_pc, rc, ia_rad)
            x1, x2, w13, L, Y, w15, F, S, K, W, tl, lc, cl, sl, ria, C, B = cb_params

            o_pc_bearing = calculate_o_pc_bearing(h1, theta_pc, w13)
            sp_pc_bearing = calculate_sp_pc_bearing(w13, h1, math.degrees(theta_pc))
            sp_xy, pc_xy, cp_xy, ps_xy = calculate_spiral_coordinates(
                bp_xy, ip_xy, ep_xy, h1, h2, tl, x1, w15, sp_pc_bearing, rc, lc, w13
            )
            center_xy = (
                math.cos(math.radians(o_pc_bearing + 180)) * rc + pc_xy[0],
                math.sin(math.radians(o_pc_bearing + 180)) * rc + pc_xy[1],
            )

            if i == 0:
                ip_sta = bp_sta + ip_dist1
            else:
                bp_sta = end_sta_list[i - 1]
                ip_sta = bp_sta + ip_dist1 - tl_list[i - 1]

            sp_sta = ip_sta - tl
            pc_sta = sp_sta + L
            cp_sta = pc_sta + lc
            ps_sta = cp_sta + L

            if bp_sta > sp_sta:
                logger.warning(
                    "곡선겹침 IP%d: BP=%s, SP=%s",
                    i + 1, format_distance(bp_sta), format_distance(sp_sta),
                )

            end_sta_list.append(ps_sta)
            end_xy_list.append(ps_xy)

            curves.append(CurveParams(
                ip_no=i + 1, ip_xy=ip_xy, curve_type=curve_type, direction=curvedir,
                h1=h1, h2=h2, ia_degree=ia_deg, center_xy=center_xy, radius=rc,
                ip_distance1=ip_dist1, ip_distance2=ip_dist2, m=m, z=z, v=v,
                params=cb_params, o_pc_bearing=o_pc_bearing,
                sp_pc_bearing=sp_pc_bearing, theta_pc=theta_pc, ip_sta=ip_sta,
                sp_sta=sp_sta, pc_sta=pc_sta, cp_sta=cp_sta, ps_sta=ps_sta,
                sp_xy=sp_xy, pc_xy=pc_xy, cp_xy=cp_xy, ps_xy=ps_xy,
            ))

        tl_list.append(tl)

        # 변경 후
        prev_xy = end_xy_list[i - 1] if i > 0 else bp_xy
        end_sta = _compute_end_sta(
            curves[-1],
            base_sta=(curves[-1].ec_sta if curves[-1].curve_type == "simplecurve"
                      else curves[-1].ps_sta),
            ec_xy=(curves[-1].ec_xy or curves[-1].ps_xy),
            ep_xy=ep_xy,
            is_last_ip=is_last_ip,
        )
        station_group = build_station_group(
            curve=curves[-1],
            bp_sta=bp_sta,
            bp_xy=prev_xy,
            end_sta=end_sta,
            unit=unit,
        )
        station_groups.append([station_group])
        polycurve.extend(row.coord for row in station_group)

    return station_groups, curves, polycurve


def _compute_end_sta(curve: CurveParams, base_sta: float, ec_xy, ep_xy, is_last_ip) -> float:
    """마지막 IP면 EP까지 연장"""
    if not is_last_ip:
        return base_sta
    end_xy = ec_xy
    return base_sta + calculate_distance(end_xy[0], end_xy[1], ep_xy[0], ep_xy[1])


# ============================================================
# Civil3D / 좌표 저장
# ============================================================
def create_civil3d_xlsx(curves: list[CurveParams]) -> None:
    Config.ensure_output_dir()
    filename = Config.OUTPUT_DIR / Config.CIVIL3D_FILENAME
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Sheet1"

    headers = ["Previous Entity ID", "Next Entity ID", "Radius", "Spiral Length"]
    for col, header in enumerate(headers, 1):
        sheet.cell(row=1, column=col, value=header)

    for i, curve in enumerate(curves, 2):
        sheet[f"A{i}"] = curve.ip_no
        sheet[f"B{i}"] = curve.ip_no + 1
        sheet[f"C{i}"] = curve.radius
        sheet[f"D{i}"] = curve.spiral_length if curve.curve_type == "cubic" else ""

    save_with_retry(wb, filename)


def save_total_coordinates(
    station_groups: list[list[list[StationRow]]], filename: Path = None
) -> list[tuple]:
    filename = filename or (Config.ensure_output_dir() / Config.TOTAL_COORD_FILENAME)
    records = []
    for group in station_groups:
        for subgroup in group:
            for row in subgroup:
                records.append((row.sta, row.coord[0], row.coord[1], row.label))

    records.sort(key=lambda r: r[0])
    unique, seen = [], set()
    for rec in records:
        key = round(rec[0], 6)
        if key in seen:
            continue
        seen.add(key)
        unique.append(rec)

    with open(filename, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["STA", "X", "Y", "LABEL"])
        for sta, x, y, label in unique:
            writer.writerow([f"{sta:.4f}", f"{x:.4f}", f"{y:.4f}", label])

    logger.info("좌표 저장 완료: %s (%d개)", filename, len(unique))
    return unique


def create_incremental_station_report(
    station_groups: list[list[list[StationRow]]],
    unit: int = 25,
    alignment_name: str = "복선중심선",
    file_ext: str = "xlsx",
) -> None:
    if file_ext not in ("xls", "xlsx"):
        raise ValueError(f"지원하지 않는 형식: {file_ext}")

    records = []
    for group in station_groups:
        for subgroup in group:
            for row in subgroup:
                if abs(row.sta / unit - round(row.sta / unit)) > 1e-7:
                    continue
                records.append((row.sta, row.coord[0], row.coord[1], row.ta))

    records.sort(key=lambda r: r[0])
    unique, seen = [], set()
    for rec in records:
        key = round(rec[0], 6)
        if key in seen:
            continue
        seen.add(key)
        unique.append(rec)

    if not unique:
        raise ValueError("출력할 증분 측점 데이터가 없습니다.")

    path = (Config.ensure_output_dir() / Config.INCREMENTAL_REPORT_FILENAME).with_suffix(f".{file_ext}")
    wb = _create_xls_workbook(unique, alignment_name, unit) if file_ext == "xls" \
         else _create_xlsx_workbook(unique, alignment_name, unit)
    wb.save(str(path))
    logger.info("증분 측점 보고서 저장: %s (%d개)", path, len(unique))


def _create_xlsx_workbook(records, alignment_name, unit):
    wb = Workbook()
    ws = wb.active
    ws.title = "선형 증분 측점 보고서"

    widths = {"A": 23, "B": 25, "C": 25, "D": 28}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:D1")
    ws["A1"] = "선형 증분 측점 보고서"
    ws["A1"].font = Font(size=16, bold=True)

    ws["A3"], ws["B3"] = "클라이언트:", "작성자:"
    ws["A4"], ws["B4"] = "Client", "Preparer"
    ws["A7"] = f"날짜: {datetime.now():%Y-%m-%d %H:%M:%S}"
    ws["A10"] = f"선형 이름: {alignment_name}"
    ws["A12"] = f"측점 범위: 시작: {format_station(records[0][0])}, 끝: {format_station(records[-1][0])}"
    ws["A13"] = f"측점 증분: {unit:.2f}"

    headers = ["측점", "Northing", "Easting", "접선 방향"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=15, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center")

    thin = Side(style="thin", color="D9E2F3")
    for r, (sta, x, y, ta) in enumerate(records, 16):
        ws.cell(r, 1, format_station(sta))
        ws.cell(r, 2, x).number_format = '#,##0.0000"미터"'
        ws.cell(r, 3, y).number_format = '#,##0.0000"미터"'
        ws.cell(r, 4, format_bearing(ta))
        for cell in ws[r][:4]:
            cell.border = Border(bottom=thin)

    ws.freeze_panes = "A16"
    ws.auto_filter.ref = f"A15:D{15 + len(records)}"
    return wb


def _create_xls_workbook(records, alignment_name, unit):
    """xls 저장 (xlwt 필요)"""
    try:
        import xlwt
    except ImportError as e:
        raise ImportError("xls 저장에는 'xlwt' 패키지가 필요합니다.") from e

    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("선형 증분 측점 보고서")

    title_style = xlwt.easyxf("font: bold on, height 320; align: horiz left;")
    header_style = xlwt.easyxf(
        "font: bold on, colour white; pattern: pattern solid, fore_colour ocean_blue;"
        "align: horiz center;"
    )
    num_style = xlwt.easyxf(num_format_str='#,##0.0000"미터"')

    ws.write_merge(0, 0, 0, 3, "선형 증분 측점 보고서", title_style)
    ws.write(2, 0, "클라이언트:")
    ws.write(3, 0, f"날짜: {datetime.now():%Y-%m-%d %H:%M:%S}")
    ws.write(5, 0, f"선형 이름: {alignment_name}")
    ws.write(6, 0, f"측점 범위: {format_station(records[0][0])} ~ {format_station(records[-1][0])}")
    ws.write(7, 0, f"측점 증분: {unit:.2f}")

    for col, h in enumerate(["측점", "Northing", "Easting", "접선 방향"]):
        ws.write(9, col, h, header_style)

    for i, (sta, x, y, ta) in enumerate(records, 10):
        ws.write(i, 0, format_station(sta))
        ws.write(i, 1, x, num_style)
        ws.write(i, 2, y, num_style)
        ws.write(i, 3, format_bearing(ta))
    return wb


# ============================================================
# 좌표 ↔ 측점 조회
# ============================================================
def get_coord_from_station(curves, station_groups, sta: float):
    """측점 → 좌표 (계산기 사용)"""
    for i, curve in enumerate(curves):
        group = station_groups[i] if i < len(station_groups) else []
        if not group or not group[0]:
            continue
        flat = group[0]
        bp_sta, ep_sta = flat[0].sta, flat[-1].sta
        if not (bp_sta <= sta <= ep_sta):
            continue

        bp_xy = flat[0].coord  # 이전 IP 끝점 or BP

        if curve.curve_type == "simplecurve":
            offset = calculate_station_offset(
                sta, curve.bc_sta, curve.bc_sta,
                curve.ec_sta, curve.ec_sta, bp_sta,
            )
            shiftx = shifty = 0.0
        else:
            offset = calculate_station_offset(
                sta, curve.sp_sta, curve.pc_sta,
                curve.cp_sta, curve.ps_sta, bp_sta,
            )
            shiftx = cal_spiral_shiftx(
                sta, curve.sp_sta, curve.pc_sta, curve.cp_sta, curve.ps_sta,
                offset, curve.params[2], curve.params[3],
            )
            shifty = cal_spiral_shifty(
                sta, curve.sp_sta, curve.pc_sta, curve.cp_sta, curve.ps_sta,
                shiftx, curve.params[2], curve.params[0],
            )

        geometry = build_geometry_from_curve(curve, bp_xy=bp_xy, shiftx=shiftx, shifty=shifty)
        calc = create_coordinate_calculator(curve)
        return calc.compute(sta, offset, geometry).as_tuple()
    return None


def perpendicular_point_on_curve(p, curve):
    px, py = p
    for i in range(len(curve) - 1):
        x1, y1 = curve[i]
        x2, y2 = curve[i + 1]
        line_vec = np.array([x2 - x1, y2 - y1])
        point_vec = np.array([px - x1, py - y1])
        line_len_sq = float(np.dot(line_vec, line_vec))
        if line_len_sq == 0:
            continue
        t = float(np.dot(point_vec, line_vec)) / line_len_sq
        if 0 <= t <= 1:
            return (x1 + t * line_vec[0], y1 + t * line_vec[1]), i
    return None, None


def get_station_at_coordinates(
    p: tuple[float, float], polycurve: list[tuple[float, float]],
    station_groups,
) -> Optional[float]:
    """좌표 → 측점"""
    curve_math = [(y, x) for x, y in polycurve]
    apoint, index = perpendicular_point_on_curve(p, curve_math)
    if apoint is None:
        return None

    dest_x, dest_y = curve_math[index]
    dest_x_i, dest_y_i = int(dest_x), int(dest_y)

    for group in station_groups:
        for subgroup in group:
            for row in subgroup:
                cy, cx = row.coord
                if int(cx) == dest_x_i and int(cy) == dest_y_i:
                    return row.sta + calculate_distance(cx, cy, apoint[0], apoint[1])
    return None


# ============================================================
# GUI
# ============================================================
class AlignmentApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("Railway Alignment Tool v10")
        self.context = AppContext()
        self._build_ui()

    def _build_ui(self):
        r = self.root

        tk.Label(r, text="계산 간격 (기본 20)").grid(row=0, column=0)
        self.entry_unit = tk.Entry(r)
        self.entry_unit.grid(row=0, column=1)

        tk.Label(r, text="시작 측점 (기본 0)").grid(row=1, column=0)
        self.entry_start = tk.Entry(r)
        self.entry_start.grid(row=1, column=1)

        tk.Label(r, text="설계 최고 속도 (기본 250)").grid(row=2, column=0)
        self.entry_speed = tk.Entry(r)
        self.entry_speed.grid(row=2, column=1)

        tk.Label(r, text="도상 유형").grid(row=3, column=0)
        self.ballast_var = tk.IntVar(value=0)
        tk.Radiobutton(r, text="자갈도상", variable=self.ballast_var, value=0).grid(row=3, column=1)
        tk.Radiobutton(r, text="콘크리트도상", variable=self.ballast_var, value=1).grid(row=3, column=2)

        tk.Button(r, text="계산 시작", command=self.run_calculation).grid(row=4, columnspan=3, pady=10)

        tk.Label(r, text="도면 축척").grid(row=5, column=0)
        self.entry_scale = tk.Entry(r)
        self.entry_scale.insert(0, "1000")
        self.entry_scale.grid(row=5, column=1)
        tk.Button(r, text="도면 출력", command=self.export_dxf).grid(row=5, column=2)

        tk.Label(r, text="측점 → 좌표").grid(row=6, column=0)
        self.entry_station = tk.Entry(r)
        self.entry_station.grid(row=6, column=1)
        tk.Button(r, text="좌표 계산", command=self.lookup_coord).grid(row=6, column=2)

        tk.Label(r, text="좌표 → 측점").grid(row=7, column=0)
        self.entry_coord = tk.Entry(r)
        self.entry_coord.grid(row=7, column=1)
        tk.Button(r, text="측점 찾기", command=self.lookup_station).grid(row=7, column=2)

        tk.Button(r, text="프로그램 종료", command=self.root.destroy).grid(row=8, columnspan=3, pady=10)

    def _read_inputs(self) -> tuple[int, float, int, str]:
        unit = int(self.entry_unit.get()) if self.entry_unit.get() else Config.DEFAULT_UNIT
        start_sta = float(self.entry_start.get()) if self.entry_start.get() else Config.DEFAULT_START_STA
        speed = int(self.entry_speed.get()) if self.entry_speed.get() else Config.DEFAULT_DESIGN_SPEED
        ballast = "자갈도상" if self.ballast_var.get() == 0 else "콘크리트도상"
        return unit, start_sta, speed, ballast

    def run_calculation(self):
        from tkinter import messagebox
        try:
            unit, start_sta, speed, ballast = self._read_inputs()
            lines = read_alignment_file()
            if not lines:
                messagebox.showerror("Error", "선형 데이터를 읽을 수 없습니다.")
                return

            alignment = parse_ip_lines(lines)
            self.context.linestring = create_linestring(alignment)
            radius_list = [p[2] for p in alignment[1:-1]]
            angles = calculate_angles(self.context.linestring)
            parameters = calculate_parameters(radius_list, speed, ballast)

            self.context.station_groups, self.context.curves, self.context.polycurve = calculate_curve(
                self.context.linestring, radius_list, angles, parameters, unit, start_sta
            )

            write_report(self.context.curves, self.context.station_groups)
            create_civil3d_xlsx(self.context.curves)
            save_total_coordinates(self.context.station_groups)
            create_incremental_station_report(self.context.station_groups, unit=unit, file_ext="xls")

            messagebox.showinfo("Success", "보고서 및 Civil3D 출력이 완료되었습니다.")
        except Exception as e:
            logger.exception("계산 중 오류")
            messagebox.showerror("Error", f"계산 실패: {e}")

    def export_dxf(self):
        from tkinter import messagebox
        try:
            scale = int(self.entry_scale.get()) if self.entry_scale.get() else Config.DEFAULT_SCALE
            if scale < Config.DEFAULT_SCALE:
                raise ValueError("축척은 1000 이상이어야 합니다.")
            if not self.context.curves:
                raise ValueError("먼저 선형 계산을 실행하세요.")
            unit, *_ = self._read_inputs()
            # 변경 후 (동일하지만 station_groups 타입이 명확해짐)
            create_dxf(
                coordinates=self.context.polycurve,
                linestring=self.context.linestring,
                curves=self.context.curves,
                station_groups=self.context.station_groups,
                unit=unit,
                filename=Config.DXF_FILENAME,
                scale=scale,
            )
            messagebox.showinfo("Success", "도면이 생성되었습니다.")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def lookup_coord(self):
        from tkinter import messagebox
        try:
            sta = float(self.entry_station.get())
            coord = get_coord_from_station(self.context.curves, self.context.station_groups, sta)
            if coord:
                messagebox.showinfo("Result", f"X = {coord[0]:.4f}, Y = {coord[1]:.4f}")
            else:
                messagebox.showerror("Error", "좌표를 찾을 수 없습니다.")
        except ValueError:
            messagebox.showerror("Input Error", "유효한 측점을 입력하세요.")

    def lookup_station(self):
        from tkinter import messagebox
        try:
            x, y = map(float, self.entry_coord.get().split(","))
            sta = get_station_at_coordinates((x, y), self.context.polycurve, self.context.station_groups)
            if sta is not None:
                messagebox.showinfo("Result", f"측점: {format_station(sta)}")
            else:
                messagebox.showerror("Error", "측점을 찾을 수 없습니다.")
        except ValueError:
            messagebox.showerror("Input Error", "좌표를 'x,y' 형식으로 입력하세요.")


def main_gui():
    root = tk.Tk()
    AlignmentApp(root)
    root.mainloop()


if __name__ == "__main__":
    main_gui()