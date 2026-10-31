import csv
from pathlib import Path

from src.railway_alignment.core.constant_config import Config
from src.railway_alignment.core.math_util import format_station, format_bearing
from src.railway_alignment.core.models import CurveParams, StationRow
from src.railway_alignment.io.report_writer import save_with_retry

from openpyxl import Workbook

import logging
import xlwt
logger = logging.getLogger(__name__)

# ============================================================
# Civil3D / 좌표 저장
# ============================================================
def create_civil3d_xlsx(curves: list[CurveParams]) -> None:
    Config.ensure_output_dir()
    filename = Config.OUTPUT_DIR / Config.CIVIL3D_FILENAME
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Sheet1"

    headers = ["Previous Entity ID", "Next Entity ID", "Radius", "Spiral Length"]
    for col, header in enumerate(headers, 1):
        sheet.cell(row=1, column=col, value=header)

    for i, curve in enumerate(curves, 2):
        sheet[f"A{i}"] = curve.ip_no
        sheet[f"B{i}"] = curve.ip_no + 1
        sheet[f"C{i}"] = curve.radius
        sheet[f"D{i}"] = curve.spiral_length if curve.curve_type == "cubic" else ""

    save_with_retry(wb, filename)


def save_total_coordinates(
    station_groups: list[list[list[StationRow]]], filename: Path = None
) -> list[tuple]:
    filename = filename or (Config.ensure_output_dir() / Config.TOTAL_COORD_FILENAME)
    records = []
    for group in station_groups:
        for subgroup in group:
            for row in subgroup:
                records.append((row.sta, row.coord[0], row.coord[1], row.label))

    records.sort(key=lambda r: r[0])
    unique, seen = [], set()
    for rec in records:
        key = round(rec[0], 6)
        if key in seen:
            continue
        seen.add(key)
        unique.append(rec)

    with open(filename, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["STA", "X", "Y", "LABEL"])
        for sta, x, y, label in unique:
            writer.writerow([f"{sta:.4f}", f"{x:.4f}", f"{y:.4f}", label])

    logger.info("좌표 저장 완료: %s (%d개)", filename, len(unique))
    return unique





def _create_xlsx_workbook(records, alignment_name, unit):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "선형 증분 측점 보고서"

    widths = {"A": 23, "B": 25, "C": 25, "D": 28}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:D1")
    ws["A1"] = "선형 증분 측점 보고서"
    ws["A1"].font = Font(size=16, bold=True)

    ws["A3"], ws["B3"] = "클라이언트:", "작성자:"
    ws["A4"], ws["B4"] = "Client", "Preparer"
    import datetime
    ws["A7"] = f"날짜: {datetime.now():%Y-%m-%d %H:%M:%S}"
    ws["A10"] = f"선형 이름: {alignment_name}"
    ws["A12"] = f"측점 범위: 시작: {format_station(records[0][0])}, 끝: {format_station(records[-1][0])}"
    ws["A13"] = f"측점 증분: {unit:.2f}"

    headers = ["측점", "Northing", "Easting", "접선 방향"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=15, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center")

    thin = Side(style="thin", color="D9E2F3")
    for r, (sta, x, y, ta) in enumerate(records, 16):
        ws.cell(r, 1, format_station(sta))
        ws.cell(r, 2, x).number_format = '#,##0.0000"미터"'
        ws.cell(r, 3, y).number_format = '#,##0.0000"미터"'
        ws.cell(r, 4, format_bearing(ta))
        for cell in ws[r][:4]:
            cell.border = Border(bottom=thin)

    ws.freeze_panes = "A16"
    ws.auto_filter.ref = f"A15:D{15 + len(records)}"
    return wb


def _create_xls_workbook(records, alignment_name, unit):
    """xls 저장 (xlwt 필요)"""
    try:
        import xlwt
    except ImportError as e:
        raise ImportError("xls 저장에는 'xlwt' 패키지가 필요합니다.") from e

    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("선형 증분 측점 보고서")

    title_style = xlwt.easyxf("font: bold on, height 320; align: horiz left;")
    header_style = xlwt.easyxf(
        "font: bold on, colour white; pattern: pattern solid, fore_colour ocean_blue;"
        "align: horiz center;"
    )
    num_style = xlwt.easyxf(num_format_str='#,##0.0000"미터"')

    ws.write_merge(0, 0, 0, 3, "선형 증분 측점 보고서", title_style)
    ws.write(2, 0, "클라이언트:")
    from datetime import datetime
    ws.write(3, 0, f"날짜: {datetime.now():%Y-%m-%d %H:%M:%S}")
    ws.write(5, 0, f"선형 이름: {alignment_name}")
    ws.write(6, 0, f"측점 범위: {format_station(records[0][0])} ~ {format_station(records[-1][0])}")
    ws.write(7, 0, f"측점 증분: {unit:.2f}")

    for col, h in enumerate(["측점", "Northing", "Easting", "접선 방향"]):
        ws.write(9, col, h, header_style)

    for i, (sta, x, y, ta) in enumerate(records, 10):
        ws.write(i, 0, format_station(sta))
        ws.write(i, 1, x, num_style)
        ws.write(i, 2, y, num_style)
        ws.write(i, 3, format_bearing(ta))
    return wb

def create_incremental_station_report(
    station_groups: list[list[list[StationRow]]],
    unit: int = 25,
    alignment_name: str = "복선중심선",
    file_ext: str = "xlsx",
) -> None:
    if file_ext not in ("xls", "xlsx"):
        raise ValueError(f"지원하지 않는 형식: {file_ext}")

    records = []
    for group in station_groups:
        for subgroup in group:
            for row in subgroup:
                if abs(row.sta / unit - round(row.sta / unit)) > 1e-7:
                    continue
                records.append((row.sta, row.coord[0], row.coord[1], row.ta))

    records.sort(key=lambda r: r[0])
    unique, seen = [], set()
    for rec in records:
        key = round(rec[0], 6)
        if key in seen:
            continue
        seen.add(key)
        unique.append(rec)

    if not unique:
        raise ValueError("출력할 증분 측점 데이터가 없습니다.")

    path = (Config.ensure_output_dir() / Config.INCREMENTAL_REPORT_FILENAME).with_suffix(f".{file_ext}")
    wb = _create_xls_workbook(unique, alignment_name, unit) if file_ext == "xls" \
         else _create_xlsx_workbook(unique, alignment_name, unit)
    wb.save(str(path))
    logger.info("증분 측점 보고서 저장: %s (%d개)", path, len(unique))


def dxf_save_with_retry(doc, filename: str, max_retries: int = None, delay: int = None) -> bool:
    """DXF 저장 재시도"""
    max_retries = max_retries or Config.SAVE_MAX_RETRIES

    delay = delay or Config.SAVE_RETRY_DELAY
    out_path = Config.ensure_output_dir() / filename
    for attempt in range(max_retries):
        try:
            doc.saveas(str(out_path))
            logger.info("DXF saved: %s", out_path)
            return True
        except PermissionError:
            logger.warning(
                "Attempt %d/%d: permission denied (%s). Retrying in %ds...",
                attempt + 1, max_retries, out_path.name, delay,
            )
            import time
            time.sleep(delay)
    logger.error("Failed to save DXF after %d attempts: %s", max_retries, out_path)
    return False